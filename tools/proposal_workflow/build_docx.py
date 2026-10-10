"""从已登记的学校样例副本生成可编辑开题审阅稿，保留未编辑包成员。"""
from __future__ import annotations

import argparse
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import re
import tempfile
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Pt, Mm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree


def font_style(style, size, east='宋体', bold=False):
    style.font.name = 'Times New Roman'
    style.font.size = Pt(size)
    style.font.bold = bold
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), east)


def styles(document):
    """段落角色取样例度量，命名样式只用于编辑与导航。"""
    roles = [('开题正文', 12, False), ('开题小节', 12, True),
             ('开题条目', 12, True), ('开题图题', 10.5, False),
             ('开题文献', 10.5, False), ('开题提纲', 12, False)]
    for name, size, bold in roles:
        s = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        s.base_style = document.styles['Normal']
        font_style(s, size, bold=bold)
        f = s.paragraph_format
        f.space_before = Pt(0); f.space_after = Pt(0)
        f.line_spacing_rule = WD_LINE_SPACING.EXACTLY; f.line_spacing = Pt(20)
        f.widow_control = True
        f.first_line_indent = Pt(24) if name == '开题正文' else Pt(0)
        f.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        f.keep_with_next = name in ('开题小节', '开题条目', '开题图题')
        if name == '开题文献':
            f.keep_together = True
        if name == '开题小节':
            level = OxmlElement('w:outlineLvl'); level.set(qn('w:val'), '1'); s._element.get_or_add_pPr().append(level)
        if name == '开题条目':
            level = OxmlElement('w:outlineLvl'); level.set(qn('w:val'), '2'); s._element.get_or_add_pPr().append(level)
        if name == '开题图题': f.alignment = WD_ALIGN_PARAGRAPH.CENTER; f.keep_with_next = False
        if name == '开题文献':
            f.first_line_indent = Pt(-20); f.left_indent = Pt(20)
    if 'Title' not in document.styles:
        document.styles.add_style('Title', WD_STYLE_TYPE.PARAGRAPH)
    title = document.styles['Title']; title.base_style = document.styles['Normal']
    font_style(title, 34, '黑体', True)


def clean_cell(cell):
    for child in list(cell._tc):
        if child.tag != qn('w:tcPr'): cell._tc.remove(child)


def plain(text):
    return re.sub(r'\*\*|`', '', text).replace('  \n', '\n').strip()


def inline(paragraph, text):
    """保留必要强调，所有正文与引用仍是可编辑文字。"""
    for i, part in enumerate(re.split(r'\*\*(.*?)\*\*', text)):
        if part:
            run = paragraph.add_run(part.replace('`', ''))
            if i % 2: run.bold = True


def render_markdown(cell, text, base):
    clean_cell(cell)
    lines = text.strip().splitlines(); i = 0; buffer = []

    def flush():
        if buffer:
            p = cell.add_paragraph(style='开题正文'); inline(p, ' '.join(buffer)); buffer.clear()

    while i < len(lines):
        line = lines[i].strip(); i += 1
        if not line: flush(); continue
        if line.startswith('|'):
            flush(); table_lines = [line]
            while i < len(lines) and lines[i].lstrip().startswith('|'):
                table_lines.append(lines[i].strip()); i += 1
            rows = [[plain(v) for v in s.strip('|').split('|')] for s in table_lines]
            rows = [row for row in rows if not all(re.fullmatch(r':?-+:?', v) for v in row)]
            if not rows: continue
            table = cell.add_table(rows=0, cols=len(rows[0])); table.autofit = False
            table._tbl.tblPr.append(OxmlElement('w:tblBorders'))
            borders = table._tbl.tblPr.find(qn('w:tblBorders'))
            for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
                b = OxmlElement('w:' + edge)
                for k,v in [('val','single'),('sz','4'),('color','000000')]: b.set(qn('w:'+k),v)
                borders.append(b)
            for n,row in enumerate(rows):
                cells = table.add_row().cells
                if n == 0: cells[0]._tc.getparent().get_or_add_trPr().append(OxmlElement('w:tblHeader'))
                for j,value in enumerate(row):
                    cells[j].width = Mm(150 / len(row)); p=cells[j].paragraphs[0]
                    p.style='开题正文'; p.paragraph_format.first_line_indent=Pt(0)
                    inline(p,value)
                    if n==0:
                        for run in p.runs: run.bold=True
            continue
        match = re.match(r'!\[([^]]+)\]\(([^)]+)\)',line)
        if match:
            flush(); image_path=(base/match[2]).resolve()
            if not image_path.is_file(): raise ValueError('图像文件缺失：'+str(image_path))
            p=cell.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next=True
            p.add_run().add_picture(str(image_path),width=Mm(150))
            docpr=p._p.xpath('.//wp:docPr')
            if docpr: docpr[0].set('descr',match[1])
            continue
        if re.match(r'^#{3,4}\s',line):
            flush(); level=len(line.split(' ')[0]);p=cell.add_paragraph(style='开题小节' if level==3 else '开题条目')
            inline(p,re.sub(r'^#+\s','',line));continue
        if re.match(r'^图[1-3]\s',line):
            flush();p=cell.add_paragraph(style='开题图题');inline(p,line);continue
        if re.match(r'^\[\d+\]\s',line):
            flush();p=cell.add_paragraph(style='开题文献');inline(p,line);continue
        if re.match(r'^(?:[　\s]*\d+\.\d+|\*\*第[一二三四五六]章)',line):
            flush();p=cell.add_paragraph(style='开题提纲');inline(p,line)
            n=re.match(r'^\s*(\d+\.\d+\.\d+)',line)
            p.paragraph_format.left_indent=Pt(24 if n else 12)
            p.paragraph_format.keep_with_next=False;continue
        if line.startswith('- '):
            flush();p=cell.add_paragraph(style='开题正文');inline(p,line[2:]);continue
        buffer.append(line)
    flush()
    if len(cell._tc)==1:cell.add_paragraph()
    # 明确覆盖样例与渲染器继承的分页属性，允许长正文单元自然跨页。
    for p in cell.paragraphs:
        p.paragraph_format.keep_together = False
        p.paragraph_format.keep_with_next = p.style is not None and p.style.name in ('开题小节', '开题条目')
        p.paragraph_format.page_break_before = False
        if p._p.xpath('.//w:drawing'):
            p.paragraph_format.keep_with_next = True
    # 表单长行的首段不能与后段绑定，否则部分Word兼容渲染器会整行移到下一页。
    cell.paragraphs[0].paragraph_format.keep_with_next = False
    for p in cell.paragraphs:
        if p.style is not None and p.style.name == '开题文献':
            p.paragraph_format.keep_together = True
    protect_groups(cell)


def protect_groups(cell):
    """用不拆分的源表格组件将已定位标题与首段、长题录保持在一起。"""
    headings = ('4. 候选复核、性质验证与执行证据', '2. 条件约束检索规约')
    paragraphs = list(cell.paragraphs)
    for n,p in enumerate(paragraphs):
        is_heading = p.text in headings
        is_reference = p.style is not None and p.style.name == '开题文献' and p.text.startswith(('[10] ', '[20] '))
        if not (is_heading or is_reference):continue
        members = [p] + ([paragraphs[n+1]] if is_heading else [])
        table = cell.add_table(rows=1,cols=1)
        added = cell._tc[-1]
        cell._tc.remove(added)
        table.autofit = False
        table.columns[0].width = cell.width
        table.cell(0,0).width = cell.width
        width = table._tbl.tblPr.find(qn('w:tblW'))
        width.set(qn('w:type'),'dxa');width.set(qn('w:w'),str(int(cell.width.twips)))
        borders = OxmlElement('w:tblBorders')
        for edge in ('top','left','bottom','right','insideH','insideV'):
            el=OxmlElement('w:'+edge);el.set(qn('w:val'),'nil');borders.append(el)
        table._tbl.tblPr.append(borders)
        margins=OxmlElement('w:tblCellMar')
        for edge in ('top','left','bottom','right'):
            el=OxmlElement('w:'+edge);el.set(qn('w:w'),'0');el.set(qn('w:type'),'dxa');margins.append(el)
        table._tbl.tblPr.append(margins)
        no_split=OxmlElement('w:cantSplit');table.rows[0]._tr.get_or_add_trPr().append(no_split)
        target=table.cell(0,0);clean_cell(target)
        p._p.addprevious(table._tbl)
        for member in members:target._tc.append(member._p)


def fill_plain(cell, text):
    p=cell.paragraphs[0]
    for child in list(p._p):
        if child.tag!=qn('w:pPr'): p._p.remove(child)
    p.add_run(text)


def build(reference, markdown, output, contract, expected_hash):
    reference,markdown,output,contract=map(lambda p:Path(p).resolve(),[reference,markdown,output,contract])
    if output==reference:raise ValueError('不能覆盖参考样例')
    actual=hashlib.sha256(reference.read_bytes()).hexdigest()
    if actual!=expected_hash or expected_hash not in contract.read_text():
        raise ValueError('参考摘要与格式契约不一致')
    text=markdown.read_text()
    title=text.splitlines()[0].lstrip('# ').strip()
    pieces=re.split(r'^## ([一二三四五])、[^\n]+\n',text,flags=re.M)
    sections=dict(zip(pieces[1::2],pieces[2::2]))
    if set(sections)!=set('一二三四五'):raise ValueError('正文五部分不完整')
    doc=Document(reference);styles(doc)
    cover=doc.paragraphs[:19]
    for n,p in enumerate(cover):
        p.paragraph_format.line_spacing_rule=WD_LINE_SPACING.EXACTLY
        p.paragraph_format.line_spacing=Pt(48 if n==0 else 44 if n==1 else 22)
        p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(2)
        if n in (2,5,6,7,8,15,17,18):
            p.text='';p.paragraph_format.line_spacing=Pt(8)
    cover[1].style='Title'
    cover[3].text='研究生类别：□学术学位博士  □专业学位博士'
    cover[4].text='            □学术学位硕士  □专业学位硕士'
    for n,label in [(9,'姓　　名'),(10,'学　　号'),(11,'专业名称'),(12,'研究方向'),(13,'导师姓名'),(14,'入学年月')]:
        p=cover[n];p.text=label+'：　[待填写]'
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:font_style(run,15)
    cover[7].text='第一版审阅稿'
    cover[7].alignment=WD_ALIGN_PARAGRAPH.CENTER
    cover[7].paragraph_format.line_spacing=Pt(22)
    cover[16].text='填表日期：　2026年10月10日（审阅稿日期）'
    cover[16].alignment=WD_ALIGN_PARAGRAPH.CENTER
    for run in cover[16].runs:font_style(run,14)
    cover[18].paragraph_format.page_break_before=True
    t=doc.tables[0]
    fill_plain(t.rows[0].cells[2],title)
    fill_plain(t.rows[1].cells[4],'[待核定]')
    fill_plain(t.rows[1].cells[9],'待核对')
    for key,row in [('一',3),('二',5),('三',7),('四',9)]:
        render_markdown(t.rows[row].cells[0],sections[key],markdown.parent)
        t.rows[row].cells[0].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
        height=t.rows[row]._tr.xpath('./w:trPr/w:trHeight')
        for element in height:element.getparent().remove(element)
    # 源样例真实导师意见必须清理，评审网格与标签仍保留。
    c=t.rows[11].cells[0];clean_cell(c)
    for value in ['指导教师审阅意见：','','','','指导教师签名：','年　月　日']:
        p=c.add_paragraph(value);p.paragraph_format.line_spacing=Pt(20)
    t.rows[10].cells[0].paragraphs[0].paragraph_format.page_break_before=True
    # 原网格及全部22行保留；移动报告会区域为同组件续表，外部段落可靠换页。
    continuation = deepcopy(t._tbl)
    for row in list(continuation.tr_lst)[:12]:continuation.remove(row)
    for row in list(t._tbl.tr_lst)[12:]:t._tbl.remove(row)
    separator=OxmlElement('w:p');props=OxmlElement('w:pPr')
    page=OxmlElement('w:pageBreakBefore');props.append(page)
    spacing=OxmlElement('w:spacing');spacing.set(qn('w:line'),'20');spacing.set(qn('w:lineRule'),'exact');props.append(spacing)
    separator.append(props);t._tbl.addnext(separator);separator.addnext(continuation)
    # Word要求表后保留段落；使用最小不可见行高，避免续表后的空段独占末页。
    tail=doc.paragraphs[-1]
    tail.paragraph_format.line_spacing_rule=WD_LINE_SPACING.EXACTLY
    tail.paragraph_format.line_spacing=Pt(1)
    tail.paragraph_format.space_before=Pt(0);tail.paragraph_format.space_after=Pt(0)
    tail.paragraph_format.keep_with_next=False;tail.paragraph_format.keep_together=False
    snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'false');tail._p.get_or_add_pPr().append(snap)
    for run in tail.runs:run.font.size=Pt(1)
    doc.core_properties.author='';doc.core_properties.last_modified_by=''
    doc.core_properties.title=title;doc.core_properties.subject='开题报告第一版审阅稿'
    settings=doc.settings._element
    update=settings.find(qn('w:updateFields'))
    if update is None:update=OxmlElement('w:updateFields');settings.append(update)
    update.set(qn('w:val'),'true')
    output.parent.mkdir(parents=True,exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='proposal-docx-') as temp:
        edited=Path(temp)/'edited.docx';doc.save(edited)
        with ZipFile(reference) as base, ZipFile(edited) as patch:
            names=set(patch.namelist());original=set(base.namelist())
            mutable={'word/document.xml','word/styles.xml','word/settings.xml',
                     'word/_rels/document.xml.rels','docProps/core.xml','[Content_Types].xml'}
            document_xml=etree.fromstring(patch.read('word/document.xml'))
            used=set(document_xml.xpath('//@r:embed|//@r:link',namespaces={'r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}))
            relations=etree.fromstring(patch.read('word/_rels/document.xml.rels'))
            removed_media=set()
            for relation in list(relations):
                if relation.get('Type','').endswith('/image') and relation.get('Id') not in used:
                    removed_media.add('word/'+relation.get('Target'));relations.remove(relation)
            relations_data=etree.tostring(relations,xml_declaration=True,encoding='UTF-8',standalone=True)
            with ZipFile(output,'w',ZIP_DEFLATED) as out:
                for info in base.infolist():
                    if info.filename in removed_media:continue
                    data=patch.read(info.filename) if info.filename in mutable else base.read(info.filename)
                    if info.filename=='word/_rels/document.xml.rels':data=relations_data
                    if info.filename=='word/footer1.xml':
                        xml=etree.fromstring(data)
                        namespaces={'wp':'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing',
                                    'a':'http://schemas.openxmlformats.org/drawingml/2006/main','v':'urn:schemas-microsoft-com:vml'}
                        for el in xml.xpath('//wp:extent|//a:ext',namespaces=namespaces):el.set('cx','457200')
                        for el in xml.xpath('//v:rect',namespaces=namespaces):
                            el.set('style',el.get('style').replace('width:13.95pt','width:36pt').replace('margin-left:428.25pt','margin-left:406.2pt'))
                        data=etree.tostring(xml,xml_declaration=True,encoding='UTF-8',standalone=True)
                    out.writestr(info,data)
                for name in names-original:out.writestr(name,patch.read(name))
    if hashlib.sha256(reference.read_bytes()).hexdigest()!=expected_hash:raise ValueError('参考文件被改变')
    return {'输出':str(output),'参考摘要':expected_hash,'输出摘要':hashlib.sha256(output.read_bytes()).hexdigest(),
            '页面与节':len(doc.sections),'说明':'文件生成不等于页面验收；须按documents技能另行渲染和逐页检查'}


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    for name in ('reference','markdown','output','contract','reference-sha256'):parser.add_argument('--'+name,required=True)
    args=parser.parse_args()
    print(json.dumps(build(args.reference,args.markdown,args.output,args.contract,args.reference_sha256),ensure_ascii=False,indent=2))


if __name__=='__main__':main()
